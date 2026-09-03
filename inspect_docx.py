import docx

doc = docx.Document('Dentia_Clinical_Odontogram_Specification.docx')
with open('docx_summary.txt', 'w', encoding='utf-8') as f:
    f.write(f"Total Paragraphs: {len(doc.paragraphs)}\n")
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip():
            f.write(f"P[{i}]: {p.text}\n")
    f.write(f"\nTotal Tables: {len(doc.tables)}\n")
    for t_idx, t in enumerate(doc.tables):
        f.write(f"\n=== Table {t_idx} (rows: {len(t.rows)}, cols: {len(t.columns)}) ===\n")
        for r_idx, r in enumerate(t.rows[:6]):
            row_txt = [c.text.strip().replace('\n', ' ') for c in r.cells]
            f.write(f"  R[{r_idx}]: {' | '.join(row_txt)}\n")

print("Dumped docx summary to docx_summary.txt")
