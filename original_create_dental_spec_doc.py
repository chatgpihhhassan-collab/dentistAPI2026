import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_document():
    doc = docx.Document()

    # Page Margins
    sections = doc.sections
    for s in sections:
        s.top_margin = Inches(0.8)
        s.bottom_margin = Inches(0.8)
        s.left_margin = Inches(0.85)
        s.right_margin = Inches(0.85)

    # Styles
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Calibri'
    style_normal.font.size = Pt(10.5)
    style_normal.font.color.rgb = RGBColor(30, 41, 59)

    # ==================== COVER / TITLE ====================
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_p.add_run("DENTIA CLINICAL AI SUITE")
    title_run.font.name = 'Arial'
    title_run.font.size = Pt(22)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(16, 36, 75)

    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = sub_p.add_run("Comprehensive 5-Surface Zone Anatomy, 7-Category EHR Clinical Palette & AI Voice Assistant Guide")
    sub_run.font.size = Pt(12)
    sub_run.font.color.rgb = RGBColor(74, 124, 210)
    sub_run.font.bold = True

    meta_p = doc.add_paragraph()
    meta_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta_run = meta_p.add_run("Clinical EHR & Interactive Odontogram Specification Document · August 2026")
    meta_run.font.size = Pt(9.5)
    meta_run.font.color.rgb = RGBColor(100, 116, 139)
    meta_run.font.italic = True

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # ==================== 1. EXECUTIVE OVERVIEW ====================
    h1 = doc.add_paragraph()
    r = h1.add_run("1. Executive Overview & System Architecture")
    r.font.name = 'Arial'
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.color.rgb = RGBColor(16, 36, 75)

    doc.add_paragraph(
        "The Dentia Clinical AI Suite provides an advanced medical-grade odontogram and EHR dental charting platform. "
        "The system combines Three.js WebGL 3D procedural/textured arch visualization across both the Maxilla (Upper Jaw, 16 teeth) "
        "and Mandible (Lower Jaw, 16 teeth) with precise 5-surface zone anatomical breakdown for localized restoration and pathology tracking."
    )

    # ==================== 2. 5-SURFACE ZONE ANATOMY ====================
    h2 = doc.add_paragraph()
    r2 = h2.add_run("2. 5-Surface Zone Diagram (Surface-Specific Dental Charting)")
    r2.font.name = 'Arial'
    r2.font.size = Pt(14)
    r2.font.bold = True
    r2.font.color.rgb = RGBColor(16, 36, 75)

    doc.add_paragraph(
        "In modern dental practice, clinical conditions such as caries, fillings, inlays, and sealants are localized to specific surfaces "
        "rather than uniformly coloring the entire tooth crown. Every human tooth is divided into 5 standard anatomical surfaces:"
    )

    # Table of 5 Surfaces
    surf_table = doc.add_table(rows=1, cols=3)
    surf_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    surf_table.autofit = False

    hdr = surf_table.rows[0].cells
    hdr[0].width = Inches(1.2)
    hdr[1].width = Inches(2.2)
    hdr[2].width = Inches(3.3)

    headers = ["Surface Code", "Anatomical Zone", "Clinical Definition & Orientation"]
    for i, title in enumerate(headers):
        set_cell_background(hdr[i], "10244B")
        set_cell_margins(hdr[i], 120, 120, 150, 150)
        p = hdr[i].paragraphs[0]
        run = p.add_run(title)
        run.font.bold = True
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(255, 255, 255)

    surfaces_data = [
        ("O / I", "Occlusal / Incisal", "Center masticatory chewing surface for molars/premolars, or cutting edge for anterior incisors."),
        ("M", "Mesial", "Anterior proximal surface oriented toward the dental arch midline."),
        ("D", "Distal", "Posterior proximal surface oriented away from the dental arch midline."),
        ("B / F", "Buccal / Facial (Labial)", "Outer tooth surface facing the cheeks (posterior) or lips (anterior incisors & canines)."),
        ("L / P", "Lingual / Palatal", "Inner tooth surface facing the tongue (mandible) or hard palate (maxilla).")
    ]

    for row_idx, (code, name, desc) in enumerate(surfaces_data):
        row = surf_table.add_row()
        bg_hex = "F8FAFC" if row_idx % 2 == 0 else "FFFFFF"
        for c_idx, text in enumerate([code, name, desc]):
            cell = row.cells[c_idx]
            set_cell_background(cell, bg_hex)
            set_cell_margins(cell, 90, 90, 120, 120)
            p = cell.paragraphs[0]
            r = p.add_run(text)
            r.font.size = Pt(9)
            if c_idx == 0:
                r.font.bold = True
                r.font.color.rgb = RGBColor(74, 124, 210)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    doc.add_paragraph(
        "Interactive Chart Workflow: Clicking any tooth in the 3D or 2D dental odontogram opens the interactive 5-Zone Geometric Diamond Inspector. "
        "Doctors can click individual zones (e.g. O, MO, DO, MOD, MODBL) to apply restorations or pathologies precisely to those anatomical faces."
    )

    # ==================== 3. 7 GROUPED EHR CATEGORIES ====================
    h3 = doc.add_paragraph()
    r3 = h3.add_run("3. 7-Category EHR Clinical Palette & Color Representation Table")
    r3.font.name = 'Arial'
    r3.font.size = Pt(14)
    r3.font.bold = True
    r3.font.color.rgb = RGBColor(16, 36, 75)

    doc.add_paragraph(
        "To prevent cognitive overload from 30+ disparate clinical conditions, the Dentia EHR groups all dental states into 7 standardized clinical domains:"
    )

    conditions_table_data = [
        # Restorative
        ("🛠️ Restorative", "Normal / Healthy", "#F3ECE2", "Plain cream/white ivory enamel, no pathology"),
        ("🛠️ Restorative", "Filling — Amalgam", "#9AA5AB", "Solid metallic silver-grey patch on surface zone"),
        ("🛠️ Restorative", "Filling — Composite", "#3B82F6", "Tooth-colored fill with distinct blue outline"),
        ("🛠️ Restorative", "Filling — GIC", "#E8D98A", "Pale glass ionomer cement yellow tint"),
        ("🛠️ Restorative", "Crown — Metal", "#B0B0B0", "Full crown solid metallic-silver fill"),
        ("🛠️ Restorative", "Crown — PFM", "#B0B0B0 / #FFFDF8", "Porcelain-Fused-to-Metal (Grey base + white occlusal rim)"),
        ("🛠️ Restorative", "Crown — Ceramic / Zirconia", "#FFFDF8", "High-gloss translucent ivory white with shine highlight"),
        ("🛠️ Restorative", "Veneer", "#CFEFF0", "Light cyan thin aesthetic overlay on facial/labial surface"),
        ("🛠️ Restorative", "Bridge / Pontic", "#8A8A8A", "Connecting bar between abutments; dashed pontic outline"),
        ("🛠️ Restorative", "Sealant", "#D9E8A0", "Light green-yellow thin coating on occlusal pits & fissures"),
        ("🛠️ Restorative", "Post & Core", "#5A5A5A", "Dark grey vertical post inside root canal space"),

        # Endodontic
        ("⚡ Endodontic", "Root Canal Treated (RCT)", "#6B4FA0", "Purple vertical line from crown to apex + apical dot"),
        ("⚡ Endodontic", "Pulp Exposure", "#EF4444", "Bright red dot at active pulp horn exposure"),
        ("⚡ Endodontic", "Periapical Lesion / Abscess", "#3D0A0A", "Dark red-black radiolucent apical halo circle"),

        # Surgical
        ("🩺 Surgical", "Missing Tooth", "#4A231A", "Empty anatomical alveolar socket depression"),
        ("🩺 Surgical", "Extraction Indicated", "#B5122E", "Red bold 'X' cross overlay over tooth crown"),
        ("🩺 Surgical", "Extracted (History)", "#C9C9C9", "Light grey dashed silhouette + soft past X record"),
        ("🩺 Surgical", "Dental Implant", "#0E8A80", "Teal titanium fixture screw threads + abutment crown"),
        ("🩺 Surgical", "Retained Root", "#6B4029", "Brown residual root apex in bone without clinical crown"),

        # Pathology
        ("🔬 Pathology", "Caries (Decay)", "#8B5A2B → #1C110A", "Brown-to-black gradient patch on affected zone"),
        ("🔬 Pathology", "Fractured / Chipped", "#7A1626", "Dark red zigzag crack line across crown"),
        ("🔬 Pathology", "Attrition / Abrasion", "#C9A66B", "Tan/yellow flat worn facet patch on occlusal table"),
        ("🔬 Pathology", "Discoloration / Fluorosis", "#D8CBAE", "Mottled brown-white irregular patches on crown"),
        ("🔬 Pathology", "Root Resorption", "#EA8C1E", "Orange dashed notch / cervical root indent"),

        # Periodontal
        ("🩸 Periodontal", "Mobility Grade I/II/III", "#22C55E / #F59E0B / #DC2626", "Roman numeral badge (I, II, III) adjacent to tooth"),
        ("🩸 Periodontal", "Gingival Recession", "#E0665A", "Pink-red line showing apical margin level + mm label"),
        ("🩸 Periodontal", "Calculus / Plaque", "#B08D3E", "Yellow-brown supragingival deposit band at cervical line"),

        # Developmental
        ("📐 Developmental", "Impacted / Unerupted", "#2F6FED", "Blue dashed outline with angled impaction trajectory"),
        ("📐 Developmental", "Supernumerary Tooth", "#D97706", "Orange dashed extra tooth silhouette + asterisk (*) mark"),
        ("📐 Developmental", "Malposition / Rotation", "#3B82F6", "Rotated tooth orientation + blue curved arrow icon"),
        ("📐 Developmental", "Diastema (Gap)", "#6B7280", "Grey double-ended arrow (↔) + interdental mm measurement"),

        # Appliance
        ("🦷 Appliance", "Orthodontic Bracket", "#C0C0C0", "Metallic silver square bracket on facial surface center"),
        ("🦷 Appliance", "Space Maintainer", "#93C5FD", "Light blue band & loop appliance across interdental gap"),
        ("🦷 Appliance", "Denture Section", "#D8A9A0", "Grey-pink dashed arch bracket wrapping spanned ridge")
    ]

    cond_table = doc.add_table(rows=1, cols=4)
    cond_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cond_table.autofit = False

    chdr = cond_table.rows[0].cells
    chdr[0].width = Inches(1.5)
    chdr[1].width = Inches(2.0)
    chdr[2].width = Inches(1.2)
    chdr[3].width = Inches(2.0)

    for i, title in enumerate(["Category", "Clinical Condition", "Color / Hex", "Visual EHR Representation"]):
        set_cell_background(chdr[i], "10244B")
        set_cell_margins(chdr[i], 120, 120, 120, 120)
        p = chdr[i].paragraphs[0]
        run = p.add_run(title)
        run.font.bold = True
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(255, 255, 255)

    for row_idx, (cat, cond, color_code, repr_desc) in enumerate(conditions_table_data):
        row = cond_table.add_row()
        bg_hex = "F8FAFC" if row_idx % 2 == 0 else "FFFFFF"
        for c_idx, text in enumerate([cat, cond, color_code, repr_desc]):
            cell = row.cells[c_idx]
            set_cell_background(cell, bg_hex)
            set_cell_margins(cell, 70, 70, 100, 100)
            p = cell.paragraphs[0]
            r = p.add_run(text)
            r.font.size = Pt(8.5)
            if c_idx == 0:
                r.font.bold = True
            elif c_idx == 1:
                r.font.bold = True
                r.font.color.rgb = RGBColor(16, 36, 75)
            elif c_idx == 2:
                r.font.bold = True
                r.font.color.rgb = RGBColor(74, 124, 210)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # ==================== 4. AI VOICE ASSISTANT GUIDE ====================
    h4 = doc.add_paragraph()
    r4 = h4.add_run("4. AI Dental Voice Assistant & Real-Time Dictation Guide")
    r4.font.name = 'Arial'
    r4.font.size = Pt(14)
    r4.font.bold = True
    r4.font.color.rgb = RGBColor(16, 36, 75)

    doc.add_paragraph(
        "The AI Dental Voice Assistant allows clinicians to dictate examination findings hands-free. "
        "The Natural Language Processing (NLP) engine parses tooth numbering (Universal 1–32 or FDI notation), "
        "surface codes (MODBL), and multi-condition diagnoses in real-time."
    )

    ai_table = doc.add_table(rows=1, cols=3)
    ai_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    ai_table.autofit = False

    ai_hdr = ai_table.rows[0].cells
    ai_hdr[0].width = Inches(2.2)
    ai_hdr[1].width = Inches(2.8)
    ai_hdr[2].width = Inches(1.7)

    for i, title in enumerate(["Doctor Dictation (Voice / Text)", "AI Assistant Automated Action", "Surface Zones Targeted"]):
        set_cell_background(ai_hdr[i], "10244B")
        set_cell_margins(ai_hdr[i], 120, 120, 120, 120)
        p = ai_hdr[i].paragraphs[0]
        run = p.add_run(title)
        run.font.bold = True
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(255, 255, 255)

    ai_examples = [
        ('"Tooth 14 MOD caries"', "Assigns active Caries (Decay) gradient (#8B5A2B) to Mesial, Occlusal, and Distal surfaces.", "M, O, D"),
        ('"Tooth 30 occlusal amalgam filling"', "Renders metallic silver-grey (#9AA5AB) amalgam restoration on Occlusal table.", "O"),
        ('"Tooth 19 root canal with periapical lesion"', "Marks Tooth #19 as Root Canal Treated (#6B4FA0) + generates apical radiolucent abscess halo.", "Whole Tooth (Endo)"),
        ('"Tooth 3 missing, tooth 4 dental implant"', "Renders Tooth #3 as empty alveolar socket (#4A231A) and Tooth #4 as titanium implant fixture (#0E8A80).", "Tooth 3 (Surgical)\nTooth 4 (Implant)"),
        ('"Tooth 8 ceramic zirconia crown and 9 veneer"', "Renders Tooth #8 with full translucent ivory zirconia crown (#FFFDF8) and Tooth #9 with light cyan facial veneer shell (#CFEFF0).", "Tooth 8 (Crown)\nTooth 9 (Facial/B)"),
        ('"Tooth 24 grade 2 mobility and 3mm recession"', "Assigns Roman numeral 'II' mobility badge (#F59E0B) and marks 3mm gingival margin recession (#E0665A).", "Tooth 24 (Perio)"),
        ('"Tooth 11 and 21 diastema 2mm"', "Places interdental double-arrow (↔) with '2.0mm' label across central incisor midline.", "Interdental Midline")
    ]

    for row_idx, (dictation, action, target) in enumerate(ai_examples):
        row = ai_table.add_row()
        bg_hex = "F8FAFC" if row_idx % 2 == 0 else "FFFFFF"
        for c_idx, text in enumerate([dictation, action, target]):
            cell = row.cells[c_idx]
            set_cell_background(cell, bg_hex)
            set_cell_margins(cell, 80, 80, 100, 100)
            p = cell.paragraphs[0]
            r = p.add_run(text)
            r.font.size = Pt(8.5)
            if c_idx == 0:
                r.font.bold = True
                r.font.color.rgb = RGBColor(16, 36, 75)
            elif c_idx == 2:
                r.font.bold = True
                r.font.color.rgb = RGBColor(74, 124, 210)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # ==================== 5. CODEBASE COMPONENT ARCHITECTURE ====================
    h5 = doc.add_paragraph()
    r5 = h5.add_run("5. Codebase Component Architecture")
    r5.font.name = 'Arial'
    r5.font.size = Pt(14)
    r5.font.bold = True
    r5.font.color.rgb = RGBColor(16, 36, 75)

    comp_items = [
        ("ToothSurfaceDiagram.jsx", "Interactive 5-zone SVG geometric cross diagram. Supports dynamic orientation (Buccal/Lingual, Mesial/Distal) based on Maxilla vs Mandible quadrants. Live multi-zone state management (MOD, DO, MO, B, L)."),
        ("ClinicalConditionPalette.jsx", "Grouped 7-category EHR condition palette with color swatches, active condition selector, and expandable AI voice commands drawer."),
        ("ThreeDentalJawArch.jsx", "WebGL 3D dental arch engine with orthographic camera, directional studio lighting, raycasting hover/click, and dynamic clinical texture/material shaders seated inside millimeter-calibrated empty jaw sockets."),
        ("ChartPage.jsx", "Master clinical odontogram orchestrator combining 3D WebGL arches, 2D dual-jaw visualizer, live audio waveform equalizer, AI SOAP note generator, and interactive patient EHR records.")
    ]

    for filename, desc in comp_items:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        run_fn = p.add_run(f"• {filename}: ")
        run_fn.font.bold = True
        run_fn.font.color.rgb = RGBColor(74, 124, 210)
        run_desc = p.add_run(desc)
        run_desc.font.size = Pt(9.5)

    # Save Document
    out_path = r"f:\DentistApp_Theme2\Dentia_Clinical_Odontogram_Specification.docx"
    doc.save(out_path)
    print(f"Document successfully created at: {out_path}")

if __name__ == '__main__':
    create_document()
